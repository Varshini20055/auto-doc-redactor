"""
Extract text (plus location info needed for redaction) from PDF and DOCX files.

For PDFs:
  - Text-based PDFs: PyMuPDF text extraction (fast, preserves layout).
  - Scanned/image PDFs: OCR via PaddleOCR (primary) or Tesseract (fallback),
    with page number, line number, and bounding box preserved.

For DOCX we keep text grouped *per paragraph*, covering body text, every
table cell (including nested tables), and every header/footer.
"""
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Iterator, Tuple
import fitz  # PyMuPDF
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table


@dataclass
class PageText:
    page_number: int   # 0-indexed
    text: str
    lines: list = field(default_factory=list)   # list of {line_number, text, bbox, source}
    extraction_source: str = "pymupdf"


@dataclass
class ParagraphText:
    key: str            # stable location key, e.g. "body:4" or "table0:r1c2p0"
    text: str


def extract_pdf_pages(path: Path) -> List[PageText]:
    """Extract text from a PDF. Auto-detects scanned PDFs and uses OCR."""
    from app.services.ocr import is_scanned_pdf, ocr_pdf

    # Only attempt OCR if the file is genuinely image-based (has page content
    # but no extractable text). A redacted PDF that had its text stripped is
    # not a scanned PDF — it just has no text left, which is correct.
    if is_scanned_pdf(path):
        try:
            ocr_pages = ocr_pdf(path)
            # If OCR returned nothing useful, fall through to PyMuPDF
            if any(op.lines for op in ocr_pages):
                result = []
                for op in ocr_pages:
                    lines_data = [
                        {"line_number": ln.line_number, "text": ln.text,
                         "bbox": ln.bbox, "source": op.source}
                        for ln in op.lines
                    ]
                    result.append(PageText(
                        page_number=op.page - 1,
                        text=op.full_text,
                        lines=lines_data,
                        extraction_source=f"ocr_{op.source}",
                    ))
                return result
        except Exception as e:
            import logging
            logging.getLogger(__name__).warning("OCR failed, falling back to PyMuPDF: %s", e)

    # Normal text-based PDF (or OCR produced no output)
    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            raw_text = page.get_text("text")
            lines_data = [
                {"line_number": j + 1, "text": ln, "bbox": [], "source": "pymupdf"}
                for j, ln in enumerate(raw_text.splitlines()) if ln.strip()
            ]
            pages.append(PageText(
                page_number=i,
                text=raw_text,
                lines=lines_data,
                extraction_source="pymupdf",
            ))
    return pages


def _iter_table_paragraphs(table: Table, prefix: str) -> Iterator[Tuple[str, Paragraph]]:
    """Yield (key, paragraph) for every paragraph in every cell of a table,
    recursing into any tables nested inside a cell."""
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            for pi, p in enumerate(cell.paragraphs):
                yield (f"{prefix}:r{ri}c{ci}p{pi}", p)
            for nti, ntable in enumerate(cell.tables):
                yield from _iter_table_paragraphs(ntable, f"{prefix}:r{ri}c{ci}nt{nti}")


def _iter_header_footer_paragraphs(document: docx.Document) -> Iterator[Tuple[str, Paragraph]]:
    """Headers/footers can carry letterhead contacts, case numbers, or patient
    names. Only emit parts that are actually defined (not inherited from a
    previous section) to avoid touching shared/linked parts redundantly."""
    for si, section in enumerate(document.sections):
        parts = [
            ("header", section.header),
            ("footer", section.footer),
            ("firstHeader", section.first_page_header),
            ("firstFooter", section.first_page_footer),
            ("evenHeader", section.even_page_header),
            ("evenFooter", section.even_page_footer),
        ]
        for label, part in parts:
            if part is None or part.is_linked_to_previous:
                continue
            for pi, p in enumerate(part.paragraphs):
                yield (f"sec{si}{label}:p{pi}", p)
            for ti, table in enumerate(part.tables):
                yield from _iter_table_paragraphs(table, f"sec{si}{label}:t{ti}")


def iter_all_docx_paragraphs(document: docx.Document) -> Iterator[Tuple[str, Paragraph]]:
    """Every paragraph in the document that could plausibly carry PII: body
    text, table cells (recursively), and headers/footers."""
    for i, p in enumerate(document.paragraphs):
        yield (f"body:{i}", p)
    for ti, table in enumerate(document.tables):
        yield from _iter_table_paragraphs(table, f"table{ti}")
    yield from _iter_header_footer_paragraphs(document)


def extract_docx_paragraphs(path: Path) -> List[ParagraphText]:
    document = docx.Document(str(path))
    return [ParagraphText(key=key, text=p.text) for key, p in iter_all_docx_paragraphs(document)]


def extract_text_any(path: Path, file_type: str) -> str:
    """Convenience: full-document plain text, e.g. for an LLM context pass."""
    if file_type == "pdf":
        return "\n".join(p.text for p in extract_pdf_pages(path))
    elif file_type == "docx":
        return "\n".join(p.text for p in extract_docx_paragraphs(path))
    raise ValueError(f"Unsupported file_type: {file_type}")


def find_table_label_findings(document: docx.Document):
    """
    Scan every top-level table for two common structured-data shapes and
    return synthetic, high-confidence findings for cells that isolated NER
    would likely miss:

      Pattern A — label/value rows:  "SSN:" | "123-45-6789"
      Pattern B — header/column:     "Name" | "Email"      (row 0)
                                      "Jane" | "jane@x.com" (row 1+)

    Returns {paragraph_key: [Finding, ...]} using the same key scheme as
    iter_all_docx_paragraphs, so results merge directly into the normal
    detection/redaction flow. Only a cell's first paragraph is targeted,
    which covers the overwhelming majority of real table cells.
    """
    from app.services.pii_detector import Finding, match_table_label  # local import: avoids a cycle at module load

    results: dict = {}

    def add(key: str, cell_text: str, entity_type: str):
        if not cell_text.strip():
            return
        results.setdefault(key, []).append(Finding(
            entity_type=entity_type, start=0, end=len(cell_text),
            score=0.90, text=cell_text, source="table_label_heuristic",
        ))

    for ti, table in enumerate(document.tables):
        prefix = f"table{ti}"

        # Determine header-row shape (Pattern B) first: if row 0 reads as a
        # run of recognized column labels, it's a header row, not a
        # label/value row — and must be excluded from Pattern A below, or a
        # header row like "Name | Email | Phone" gets its own "Email"/"Phone"
        # cells wrongly tagged PERSON just because column 0 said "Name".
        is_header_table = False
        col_entities = []
        if len(table.rows) >= 2:
            header_cells = table.rows[0].cells
            col_entities = [match_table_label(c.text) for c in header_cells]
            recognized = sum(1 for e in col_entities if e)
            is_header_table = recognized >= 2 and recognized >= len(col_entities) / 2

        # Pattern A: label in column 0, value(s) in the rest of the row.
        for ri, row in enumerate(table.rows):
            if ri == 0 and is_header_table:
                continue  # this row is a column-header row, not a label/value row
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text
            entity_type = match_table_label(label)
            if not entity_type:
                continue
            for ci in range(1, len(row.cells)):
                cell = row.cells[ci]
                if not cell.paragraphs:
                    continue
                key = f"{prefix}:r{ri}c{ci}p0"
                add(key, cell.paragraphs[0].text, entity_type)

        # Pattern B: header row of recognized labels, data in the rows below.
        if is_header_table:
            for ri in range(1, len(table.rows)):
                row = table.rows[ri]
                for ci, cell in enumerate(row.cells):
                    if ci >= len(col_entities) or not col_entities[ci]:
                        continue
                    if not cell.paragraphs:
                        continue
                    key = f"{prefix}:r{ri}c{ci}p0"
                    add(key, cell.paragraphs[0].text, col_entities[ci])

    return results
